"""Document parser for extracting text from various file formats."""

import logging
import io
from typing import Optional, Dict
from .xlsx_optimizer import XLSXOptimizer

logger = logging.getLogger(__name__)


class DocumentParser:
    """
    Parser for extracting text content from uploaded documents.
    
    Supported formats:
    - PDF (.pdf)
    - Word Documents (.docx)
    - PowerPoint (.pptx)
    - Excel (.xlsx) - with intelligent optimization for attendance sheets
    - Plain Text (.txt)
    - Markdown (.md)
    """
    
    def __init__(self, optimize_xlsx: bool = True, enable_delta_detection: bool = True):
        """
        Initialize document parser.
        
        Args:
            optimize_xlsx: Enable smart XLSX processing (reduces token usage)
            enable_delta_detection: Enable delta detection for similar documents
        """
        self.optimize_xlsx = optimize_xlsx
        self.xlsx_optimizer = XLSXOptimizer(enable_delta_detection=enable_delta_detection) if optimize_xlsx else None
    
    def extract_content(self, content: bytes, extension: str, url: Optional[str] = None) -> str:
        """
        Extract text content from a document (legacy method for backwards compatibility).
        
        Args:
            content: Raw file bytes
            extension: File extension (with or without dot)
            url: Optional source URL (for delta detection)
            
        Returns:
            Extracted text content as string
        """
        try:
            # Normalize extension
            ext = extension.lstrip('.')
            filename = f"document.{ext}"
            
            # Use parse method and extract text
            result = self.parse(content, filename, url)
            
            # For XLSX with summary, prefer summary for LLM context
            if result.get('is_attendance') and result.get('summary'):
                logger.info(f"Using optimized summary for attendance XLSX (token reduction)")
                return result['summary']
            
            return result.get('text', '')
        except Exception as e:
            logger.error(f"Error in extract_content for {extension}: {e}", exc_info=True)
            # Return empty string rather than raising, to prevent redelivery loops
            return ""
    
    def parse(self, content: bytes, filename: str, url: Optional[str] = None) -> Dict[str, any]:
        """
        Extract text content from a document.
        
        Args:
            content: Raw file bytes
            filename: Original filename (used to determine format)
            url: Optional source URL (for delta detection)
            
        Returns:
            Dict with 'text' and optional 'metadata', 'summary', 'delta'
        """
        ext = filename.lower().split('.')[-1] if '.' in filename else ''
        
        try:
            if ext == 'pdf':
                text = self._parse_pdf(content)
                return {'text': text, 'metadata': {'type': 'pdf'}}
            elif ext == 'docx':
                text = self._parse_docx(content)
                return {'text': text, 'metadata': {'type': 'docx'}}
            elif ext == 'pptx':
                text = self._parse_pptx(content)
                return {'text': text, 'metadata': {'type': 'pptx'}}
            elif ext == 'xlsx':
                if self.optimize_xlsx:
                    # Use optimized XLSX processing
                    return self.xlsx_optimizer.parse_smart(content, filename, url)
                else:
                    # Fallback to standard extraction
                    text = self._parse_xlsx(content)
                    return {'text': text, 'metadata': {'type': 'xlsx'}}
            elif ext in ('txt', 'md'):
                text = self._parse_text(content)
                return {'text': text, 'metadata': {'type': ext}}
            else:
                logger.warning(f"Unknown file extension: {ext}, treating as text")
                text = self._parse_text(content)
                return {'text': text, 'metadata': {'type': 'unknown'}}
        except Exception as e:
            logger.error(f"Error parsing {filename}: {e}")
            raise ValueError(f"Failed to parse document: {e}")
    
    def _parse_pdf(self, content: bytes) -> str:
        """Extract text from PDF."""
        try:
            from pypdf import PdfReader
            
            pdf_file = io.BytesIO(content)
            reader = PdfReader(pdf_file)
            
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            
            return '\n\n'.join(text_parts)
        except ImportError:
            logger.error("pypdf not installed. Install with: pip install pypdf")
            raise
    
    def _parse_docx(self, content: bytes) -> str:
        """Extract text from Word document."""
        try:
            from docx import Document
            
            doc_file = io.BytesIO(content)
            doc = Document(doc_file)
            
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_parts.append(row_text)
            
            return '\n\n'.join(text_parts)
        except ImportError:
            logger.error("python-docx not installed. Install with: pip install python-docx")
            raise
    
    def _parse_pptx(self, content: bytes) -> str:
        """Extract text from PowerPoint presentation."""
        try:
            from pptx import Presentation
            
            pptx_file = io.BytesIO(content)
            prs = Presentation(pptx_file)
            
            text_parts = []
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_texts = [f"--- Slide {slide_num} ---"]
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text)
                
                if len(slide_texts) > 1:  # More than just the slide header
                    text_parts.append('\n'.join(slide_texts))
            
            return '\n\n'.join(text_parts)
        except ImportError:
            logger.error("python-pptx not installed. Install with: pip install python-pptx")
            raise
    
    def _parse_xlsx(self, content: bytes) -> str:
        """Extract text from Excel spreadsheet."""
        try:
            from openpyxl import load_workbook
            
            xlsx_file = io.BytesIO(content)
            wb = load_workbook(xlsx_file, read_only=True, data_only=True)
            
            text_parts = []
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_texts = [f"--- Sheet: {sheet_name} ---"]
                
                for row in sheet.iter_rows(values_only=True):
                    row_values = [str(cell) if cell is not None else '' for cell in row]
                    row_text = ' | '.join(v for v in row_values if v)
                    if row_text:
                        sheet_texts.append(row_text)
                
                if len(sheet_texts) > 1:
                    text_parts.append('\n'.join(sheet_texts))
            
            wb.close()
            return '\n\n'.join(text_parts)
        except ImportError:
            logger.error("openpyxl not installed. Install with: pip install openpyxl")
            raise
    
    def _parse_text(self, content: bytes) -> str:
        """Parse plain text or markdown files."""
        # Try common encodings
        encodings = ['utf-8', 'utf-16', 'latin-1', 'cp1252']
        
        for encoding in encodings:
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        
        # Fallback: decode with replacement
        return content.decode('utf-8', errors='replace')
